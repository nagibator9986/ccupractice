"""Excel / CSV import of practice students, group by group.

Pipeline (stateless — the SPA re-sends the file for the commit step, so there is
no server-side temp storage to lose between gunicorn workers):

    bytes ──parse_file──▶ ParsedBook ──build_plan──▶ Plan ──apply_plan──▶ DB

``build_plan`` is the ONLY place that decides what happens to a row. The preview
endpoint returns the plan; the commit endpoint rebuilds the very same plan and
applies it, so what the operator confirmed is what gets written.

Design rules
* Add-only by default. An existing student is matched by ИИН (authoritative),
  then by ФИО (+ group); a match is never duplicated. ``mode`` decides whether a
  match is left alone (``skip``), only has its EMPTY fields filled (``fill``,
  default — never clobbers hand-entered data) or is updated from the file
  (``overwrite``).
* ``is_grant_student`` is only ever RAISED, never cleared (clearing is guarded by
  the LMS invariants in ``api/students.py``).
* Partners are never auto-created from a spreadsheet — a free-text company name
  is not a legal entity. An unknown name is a warning and the field is skipped.
* Problems that make a row unsafe are ``errors`` (row skipped); recoverable
  oddities are ``warnings`` (row still imported).
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

from ..extensions import db
from ..models import Contract, Partner, Specialty, Student
from ..utils.iin import birth_date_from_iin, iin_checksum_ok, normalize_iin
from ..utils.time import utc_today
from .xlsx_utils import autosize, style_header, table_sheet

MAX_ROWS = 5000
MAX_COLUMNS = 60
_HEADER_SCAN_ROWS = 25
_BLANK_RUN_LIMIT = 300  # stop scanning a sheet after this many empty rows

MODES = ("skip", "fill", "overwrite")
MODE_LABELS = {
    "skip": "Только новые (существующих не трогать)",
    "fill": "Дополнить пустые поля существующих",
    "overwrite": "Обновить существующих данными из файла",
}


class ImportFileError(ValueError):
    """The uploaded file cannot be read as a student list (user-facing text)."""


# ─────────────────────────────────────────────────────────────────────────────
# Field catalogue
# ─────────────────────────────────────────────────────────────────────────────

@dataclass(frozen=True)
class FieldDef:
    key: str                 # parsed-row key
    label: str               # header written to the template / export
    kind: str                # name|iin|group|str|int|date|bool|partner
    attr: str | None = None  # Student attribute (defaults to key)
    max_len: int | None = None
    required: bool = False
    hint: str = ""

    @property
    def student_attr(self) -> str:
        return self.attr or self.key


FIELDS: tuple[FieldDef, ...] = (
    FieldDef("full_name", "ФИО", "name", max_len=200, required=True,
             hint="Фамилия Имя Отчество полностью"),
    FieldDef("iin", "ИИН", "iin", max_len=20, required=True,
             hint="12 цифр. В Excel задан текстовый формат, чтобы не терялся ведущий 0"),
    FieldDef("group_name", "Группа", "group", max_len=60, required=True,
             hint="Например: ПО 25-1. Можно не заполнять, если группа = название листа"),
    FieldDef("specialty", "Специальность", "str", max_len=200),
    FieldDef("specialty_code", "Код специальности", "str", max_len=60,
             hint="Если пусто — подставится из справочника «Данные» по названию специальности"),
    FieldDef("course", "Курс", "int", hint="Число от 1 до 6"),
    FieldDef("phone", "Телефон", "str", max_len=60),
    FieldDef("birth_date", "Дата рождения", "date",
             hint="ДД.ММ.ГГГГ. Если пусто — берётся из ИИН"),
    FieldDef("partner", "Предприятие (база практики)", "partner", attr="partner_id",
             hint="Название или БИН из реестра «Партнёры»"),
    FieldDef("practice_start", "Начало практики", "date", hint="ДД.ММ.ГГГГ"),
    FieldDef("practice_end", "Окончание практики", "date", hint="ДД.ММ.ГГГГ"),
    FieldDef("college_supervisor", "Руководитель практики от колледжа", "str", max_len=200),
    FieldDef("partner_supervisor", "Руководитель практики от предприятия", "str", max_len=200),
    FieldDef("practice_type", "Вид практики", "str", max_len=120,
             hint="учебной / профессиональной / производственной / преддипломной"),
    FieldDef("form_of_study", "Форма обучения", "str", max_len=60),
    FieldDef("education_program", "Образовательная программа", "str", max_len=300),
    FieldDef("enrollment_year", "Год поступления", "int"),
    FieldDef("id_card_number", "№ удостоверения личности", "str", max_len=60),
    FieldDef("id_card_issued_by", "Кем и когда выдано", "str", max_len=200),
    FieldDef("home_address", "Домашний адрес", "str", max_len=300),
    FieldDef("legal_rep_full_name", "ФИО законного представителя", "str", max_len=200),
    FieldDef("legal_rep_iin", "ИИН законного представителя", "iin", max_len=20),
    FieldDef("legal_rep_phone", "Телефон законного представителя", "str", max_len=60),
    FieldDef("is_grant_student", "Грант (госзаказ)", "bool", hint="Да / Нет"),
    FieldDef("notes", "Примечание", "str"),
)
FIELD_BY_KEY = {f.key: f for f in FIELDS}

# Pseudo-keys: columns that exist in real-world sheets but are not stored 1:1.
_NAME_PARTS = ("last_name", "first_name", "middle_name")
_IGNORED_KEYS = {"row_no", "qualification", "language"}

# Fields an operator can set once for the whole upload ("practice settings of
# the group"). Everything else only comes from the file.
DEFAULT_KEYS = (
    "partner_id", "practice_start", "practice_end", "college_supervisor",
    "partner_supervisor", "practice_type", "course", "specialty",
    "specialty_code", "form_of_study", "education_program", "enrollment_year",
)

# Empty-looking sheet names that must NOT be taken for a group name.
_GENERIC_SHEET = re.compile(
    r"^(лист|sheet|студенты|студент|список|списки|шаблон|данные|students|table|таблица|"
    r"инструкция|справочники?)\s*[\d_ -]*$",
    re.IGNORECASE,
)
_TITLE_GROUP = re.compile(
    r"групп[аыуе]?\s*[:№#\-–—]?\s*([A-Za-zА-Яа-яЁёӘәҒғҚқҢңӨөҰұҮүҺһІі][\w\-–— ./]{0,24})",
    re.IGNORECASE,
)


def _norm_header(value: Any) -> str:
    s = str(value or "").replace(" ", " ").casefold().replace("ё", "е")
    s = re.sub(r"[^\w]+", " ", s, flags=re.UNICODE)
    return re.sub(r"\s+", " ", s).strip()


def match_field(header: Any) -> str | None:
    """Map a spreadsheet header to a field key (or an ignored/pseudo key).

    Exact aliases first, then keyword rules ordered from most to least
    specific — real files say «ФИО студента», «Тел. родителя», «Практика с»…
    """
    h = _norm_header(header)
    if str(header or "").strip() in {"№", "#", "№№", "N°", "№ п/п", "№п/п"}:
        return "row_no"
    if not h:
        return None
    if h in {"п п", "пп", "n", "no", "номер", "n п п", "номер п п", "порядковый номер", "id"}:
        return "row_no"

    rep = any(w in h for w in ("родител", "представител", "опекун"))
    padded = f" {h} "

    if "иин" in h.split() or h in {"жсн", "иин студента", "иин обучающегося"} or h.startswith("иин "):
        return "legal_rep_iin" if rep else "iin"
    if "телефон" in h or h.startswith("тел ") or h == "тел":
        return "legal_rep_phone" if rep else "phone"
    if h.startswith("фио") or "фамилия имя отчество" in h or h in {
        "студент", "обучающийся", "учащийся", "ученик", "абитуриент", "ф и о",
    }:
        return "legal_rep_full_name" if rep else "full_name"
    if rep and ("фио" in h or h in {"родитель", "законный представитель", "представитель"}):
        return "legal_rep_full_name"
    if h == "фамилия":
        return "last_name"
    if h == "имя":
        return "first_name"
    if h == "отчество":
        return "middle_name"
    if "групп" in h:
        return "group_name"
    if "специальност" in h or h in {"направление", "мамандык"}:
        return "specialty_code" if ("код" in h or "шифр" in h) else "specialty"
    if h in {"код", "шифр"}:
        return "specialty_code"
    if "образовательн" in h and "программ" in h or h in {"оп", "программа"}:
        return "education_program"
    if "квалификац" in h:
        return "qualification"
    if "язык" in h:
        return "language"
    if h == "курс" or h.endswith(" курс") or h.startswith("курс "):
        return "course"
    if "рожден" in h:
        return None if "место" in h else "birth_date"
    if "удостоверен" in h or "документ" in h and "личн" in h:
        return "id_card_issued_by" if ("кем" in h or "выдан" in h) else "id_card_number"
    if ("кем" in h and "выдан" in h) or "орган выдачи" in h:
        return "id_card_issued_by"
    if "адрес" in h or "место жительства" in h:
        return "home_address"
    if "форма обучения" in h:
        return "form_of_study"
    if "год поступления" in h or "год набора" in h:
        return "enrollment_year"
    if "руководител" in h or "наставник" in h:
        if "колледж" in h:
            return "college_supervisor"
        if any(w in h for w in ("предприят", "организац", "наставник", "база")):
            return "partner_supervisor"
        return None
    if "практик" in h:
        if "вид" in h or "тип" in h:
            return "practice_type"
        if any(w in h for w in ("оконч", "конец", "завершен")) or h.endswith(" по"):
            return "practice_end"
        if "начал" in h or h.endswith(" с") or padded.startswith(" практика с"):
            return "practice_start"
        if any(w in h for w in ("предприят", "организац", "база", "место")):
            return "partner"
    if h in {"начало", "с"}:
        return "practice_start"
    if h in {"окончание", "конец", "по"}:
        return "practice_end"
    if h.split()[:1] == ["бин"] or h == "бин":
        return "partner_bin"
    if any(w in h for w in ("предприят", "организац", "партнер", "база практики", "место практики")):
        return "partner"
    if "грант" in h or "госзаказ" in h or "финансиров" in h:
        return "is_grant_student"
    if "примечан" in h or "комментар" in h or "заметк" in h:
        return "notes"
    return None


# ─────────────────────────────────────────────────────────────────────────────
# Cell parsing
# ─────────────────────────────────────────────────────────────────────────────

def _is_blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _text(v: Any) -> str:
    """Cell → clean single-line string (numbers without a trailing .0)."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Да" if v else "Нет"
    if isinstance(v, float):
        if v != v:
            return ""
        return str(int(v)) if v == int(v) else str(v)
    if isinstance(v, (datetime, date)):
        return v.strftime("%d.%m.%Y")
    return re.sub(r"\s+", " ", str(v).replace(" ", " ")).strip()


_DATE_FORMATS = ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d,%m,%Y", "%d %m %Y")


def _parse_date(v: Any) -> date | None:
    """Excel datetime / serial number / common text formats → date.

    Raises ``ValueError`` for a non-empty value that cannot be understood so the
    caller can warn instead of silently dropping it.
    """
    if _is_blank(v):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        n = float(v)
        if 1 <= n <= 80000:  # Excel serial day
            return (datetime(1899, 12, 30) + timedelta(days=int(n))).date()
        raise ValueError(str(v))
    s = _text(v)
    s = re.sub(r"\s*(г\.?|года?)\s*$", "", s, flags=re.IGNORECASE)
    s = re.sub(r"\s+\d{1,2}:\d{2}(:\d{2})?$", "", s)
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(s)


def _parse_int(v: Any) -> int | None:
    if _is_blank(v):
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return int(v)
    m = re.search(r"\d+", _text(v))
    if not m:
        raise ValueError(_text(v))
    return int(m.group())


_TRUE_WORDS = ("да", "yes", "y", "true", "1", "+", "грант", "госзаказ", "гос", "гос.", "бюджет", "грантник")
_FALSE_WORDS = ("нет", "no", "n", "false", "0", "-", "платно", "платное", "платн", "договор")


def _parse_bool(v: Any) -> bool | None:
    if _is_blank(v):
        return None
    if isinstance(v, bool):
        return v
    s = _text(v).casefold()
    if s in _TRUE_WORDS or s.startswith(("грант", "госзаказ", "бюджет")):
        return True
    if s in _FALSE_WORDS or s.startswith(("платн",)):
        return False
    raise ValueError(s)


def normalize_name(value: str) -> str:
    """Comparison key for a ФИО: case, «ё»/«е» and spacing drop out."""
    return " ".join((value or "").split()).casefold().replace("ё", "е")


def clean_full_name(value: str) -> str:
    s = " ".join((value or "").replace(" ", " ").split())
    if s and (s.isupper() or s.islower()):
        s = s.title()
    return s


_LATIN_TO_CYR = str.maketrans({
    "A": "А", "B": "В", "C": "С", "E": "Е", "H": "Н", "K": "К", "M": "М", "O": "О",
    "P": "Р", "T": "Т", "X": "Х", "Y": "У", "a": "а", "c": "с", "e": "е", "o": "о",
    "p": "р", "x": "х", "y": "у", "k": "к", "m": "м", "t": "т", "h": "н", "b": "в",
})


def normalize_group(value: str) -> str:
    """Trim, unify dashes/spaces; map Latin look-alikes to Cyrillic when the
    name already contains Cyrillic («ПO 25-1» typed with a Latin O is common)."""
    s = " ".join((value or "").replace(" ", " ").split())
    s = re.sub(r"\s*[–—−]\s*", "-", s)
    if re.search(r"[А-Яа-яЁёӘәҒғҚқҢңӨөҰұҮүҺһІі]", s):
        s = s.translate(_LATIN_TO_CYR)
    return s


def group_key(value: str) -> str:
    return re.sub(r"[\s\-_.]+", "", normalize_group(value)).casefold()


# ─────────────────────────────────────────────────────────────────────────────
# File → ParsedBook
# ─────────────────────────────────────────────────────────────────────────────

@dataclass
class ParsedSheet:
    name: str
    header_row: int = 0
    title_group: str = ""
    columns: list[dict] = field(default_factory=list)   # [{header, key, label}]
    ignored: list[str] = field(default_factory=list)
    rows: list[tuple[int, dict[str, Any]]] = field(default_factory=list)
    skipped_reason: str = ""
    has_group_column: bool = False
    headerless: bool = False
    # «Группа: ПО 26-1 / Специальность: 0414… / Курс: 1» block above the table.
    meta: dict[str, Any] = field(default_factory=dict)


@dataclass
class ParsedBook:
    filename: str
    sheets: list[ParsedSheet]


def _read_grid_xlsx(data: bytes) -> list[tuple[str, list[list[Any]]]]:
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    except Exception as exc:  # noqa: BLE001 — openpyxl raises many types
        raise ImportFileError(
            "Не удалось открыть файл. Нужен формат .xlsx (Excel 2007+). "
            "Файлы .xls сначала сохраните как .xlsx."
        ) from exc
    out: list[tuple[str, list[list[Any]]]] = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        grid: list[list[Any]] = []
        blank_run = 0
        for row in ws.iter_rows(values_only=True, max_col=MAX_COLUMNS):
            if len(grid) > MAX_ROWS + _HEADER_SCAN_ROWS:
                break
            cells = list(row)
            if all(_is_blank(c) for c in cells):
                blank_run += 1
                if blank_run > _BLANK_RUN_LIMIT:
                    break
            else:
                blank_run = 0
            grid.append(cells)
        # Group names are often a vertically merged cell: only the top-left cell
        # holds the value. Spread it down so every row sees its group.
        for rng in ws.merged_cells.ranges:
            if rng.max_row <= rng.min_row:
                continue
            top = ws.cell(row=rng.min_row, column=rng.min_col).value
            if top is None:
                continue
            for r in range(rng.min_row, min(rng.max_row, len(grid)) + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    if r - 1 < len(grid) and c - 1 < len(grid[r - 1]) and _is_blank(grid[r - 1][c - 1]):
                        grid[r - 1][c - 1] = top
        out.append((ws.title, grid))
    wb.close()
    return out


def _read_grid_csv(data: bytes, filename: str) -> list[tuple[str, list[list[Any]]]]:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ImportFileError("Не удалось определить кодировку CSV (нужна UTF-8 или Windows-1251).")
    sample = text[:4096]
    delimiter = max((";", ",", "\t"), key=sample.count)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    grid = [row for _, row in zip(range(MAX_ROWS + _HEADER_SCAN_ROWS), reader)]
    return [(Path(filename).stem or "CSV", grid)]


def parse_file(data: bytes, filename: str, column_overrides: dict | None = None) -> ParsedBook:
    ext = Path(filename or "").suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        grids = _read_grid_xlsx(data)
    elif ext in (".csv", ".txt"):
        grids = _read_grid_csv(data, filename)
    elif ext == ".xls":
        raise ImportFileError("Формат .xls устарел. Откройте файл в Excel и сохраните как «Книга Excel (.xlsx)».")
    else:
        raise ImportFileError("Допустимы файлы .xlsx или .csv")
    if not grids:
        raise ImportFileError("В файле нет ни одного видимого листа.")

    overrides = column_overrides or {}
    def _parse_all(allow_headerless: bool) -> list[ParsedSheet]:
        return [
            _parse_sheet(name, grid, {**(overrides.get("*") or {}), **(overrides.get(name) or {})},
                         allow_headerless=allow_headerless)
            for name, grid in grids
        ]

    sheets = _parse_all(False)
    if not any(s.rows for s in sheets):
        # Nothing has a header row. Only now fall back to «bare list of names
        # under a Группа block» — doing it earlier would turn the free-text
        # sheets of our own template (instructions, dictionaries) into students.
        sheets = _parse_all(True)
    if not any(s.rows for s in sheets):
        reasons = "; ".join(f"«{s.name}»: {s.skipped_reason}" for s in sheets if s.skipped_reason)
        raise ImportFileError(
            "В файле не найдено ни одной строки со студентами."
            + (f" {reasons}" if reasons else "")
            + " Скачайте шаблон — в нём нужные заголовки уже расставлены, "
              "или укажите, какая колонка содержит ФИО."
        )
    total = sum(len(s.rows) for s in sheets)
    if total > MAX_ROWS:
        raise ImportFileError(f"Слишком много строк ({total}). За один раз можно загрузить до {MAX_ROWS}.")
    return ParsedBook(filename=filename, sheets=sheets)


_META_LABELS = {
    "группа": "group_name", "учебная группа": "group_name", "тобы": "group_name",
    "специальность": "specialty", "мамандыгы": "specialty", "мамандығы": "specialty",
    "курс": "course", "куратор": "curator", "классный руководитель": "curator",
}
_META_LINE = re.compile(r"^\s*([^\W\d_][^\W\d_ ]*(?: [^\W\d_][^\W\d_ ]*){0,2})\s*[:：]\s*(.*)$", re.UNICODE)
_NAME_WORD = r"[^\W\d_]+(?:[\-'’.][^\W\d_]+)*\.?"
_NAME_LIKE = re.compile(rf"^{_NAME_WORD}(?:\s+{_NAME_WORD}){{1,3}}$", re.UNICODE)
_SPECIALTY_CODE = re.compile(r"^(\d{6,9})\s+(.+)$")


def _scan_meta(rows: list[list[Any]]) -> tuple[dict[str, Any], int]:
    """Key/value lines above the table → ``({field: value}, index of last meta row)``.

    Handles «Группа: МГК 26-1» in one cell as well as «Группа:» | «МГК 26-1» in
    neighbouring cells, and splits «04140100 Маркетинг» into code + name.
    """
    meta: dict[str, Any] = {}
    last = -1
    for i, row in enumerate(rows):
        for idx, cell in enumerate(row):
            if not isinstance(cell, str):
                continue
            text = _text(cell)
            m = _META_LINE.match(text)
            if m:
                label, rest = _norm_header(m.group(1)), m.group(2).strip()
            elif _norm_header(text) in _META_LABELS and len(text) < 30:
                label, rest = _norm_header(text), ""
            else:
                continue
            key = _META_LABELS.get(label)
            if not key:
                continue
            last = i
            value = rest
            if not value:
                value = next((_text(n) for n in row[idx + 1:] if not _is_blank(n)), "")
            if not value or key in meta:
                continue
            if key == "specialty":
                code = _SPECIALTY_CODE.match(value)
                if code:
                    meta["specialty_code"], value = code.group(1), code.group(2).strip()
            elif key == "group_name":
                value = normalize_group(value).strip(" .,:;")
            elif key == "course":
                try:
                    value = _parse_int(value)
                except ValueError:
                    continue
            meta[key] = value
    return meta, last


def _find_name_column(grid: list[list[Any]], start: int) -> int | None:
    """Header-less list: the column holding the most «Фамилия Имя [Отчество]» cells."""
    counts: dict[int, int] = {}
    for row in grid[start:start + 2000]:
        for j, cell in enumerate(row):
            if isinstance(cell, str) and _NAME_LIKE.match(_clean_cell_name(cell)):
                counts[j] = counts.get(j, 0) + 1
    if not counts:
        return None
    col, n = max(counts.items(), key=lambda kv: kv[1])
    return col if n >= 2 else None


def _clean_cell_name(value: str) -> str:
    return re.sub(r"^\s*\d+[.)]?\s+", "", _text(value))


def _parse_sheet(name: str, grid: list[list[Any]], overrides: dict[str, str] | None = None,
                 allow_headerless: bool = False) -> ParsedSheet:
    overrides = overrides or {}
    sheet = ParsedSheet(name=name)

    def _key_for(cell: Any) -> str | None:
        raw = _text(cell)
        if raw in overrides:
            return overrides[raw] or None
        return match_field(cell)

    header_idx = None
    mapping: dict[int, str] = {}
    for i, row in enumerate(grid[:_HEADER_SCAN_ROWS]):
        keys = {j: _key_for(c) for j, c in enumerate(row) if not _is_blank(c)}
        found = {k for k in keys.values() if k}
        if "full_name" in found or {"last_name", "first_name"} <= found:
            header_idx = i
            mapping = {j: k for j, k in keys.items() if k}
            break

    if header_idx is None:
        # No header row: a bare column of names under a «Группа: …» block is how
        # most group lists are kept in practice.
        meta, last_meta = _scan_meta(grid[:_HEADER_SCAN_ROWS])
        # A sheet qualifies only if it announces a group (block above the list or
        # a meaningful sheet name) — otherwise it is just some other table.
        qualifies = allow_headerless and (meta.get("group_name") or not _GENERIC_SHEET.match(name.strip()))
        col = _find_name_column(grid, last_meta + 1) if qualifies else None
        if col is None:
            sheet.skipped_reason = "не найдена колонка с ФИО (нет строки заголовков и списка «Фамилия Имя»)"
            sheet.meta = meta
            return sheet
        sheet.headerless = True
        sheet.meta = meta
        sheet.title_group = meta.get("group_name", "")
        sheet.header_row = 0
        sheet.columns.append({"header": "(без заголовка)", "key": "full_name", "label": "ФИО"})
        for i in range(last_meta + 1, len(grid)):
            row = grid[i]
            cell = row[col] if col < len(row) else None
            if isinstance(cell, str) and _NAME_LIKE.match(_clean_cell_name(cell)):
                sheet.rows.append((i + 1, {"full_name": _clean_cell_name(cell)}))
        return sheet

    sheet.header_row = header_idx + 1
    header_cells = grid[header_idx]
    sheet.meta, _ = _scan_meta(grid[:header_idx])
    sheet.title_group = sheet.meta.get("group_name", "")

    # «Список студентов группы ПО 25-1» (no colon) above the table names the group.
    if not sheet.title_group:
        for row in grid[:header_idx]:
            for cell in row:
                if isinstance(cell, str):
                    m = _TITLE_GROUP.search(cell)
                    if m:
                        sheet.title_group = normalize_group(m.group(1)).strip(" .,:;")
                        break
            if sheet.title_group:
                break

    seen: set[str] = set()
    key_cols: dict[str, int] = {}
    for j, key in mapping.items():
        if key in seen:
            continue  # first column wins on duplicates
        seen.add(key)
        key_cols[key] = j
        if key not in _IGNORED_KEYS:
            label = FIELD_BY_KEY[key].label if key in FIELD_BY_KEY else _PSEUDO_LABELS.get(key, key)
            sheet.columns.append({"header": _text(header_cells[j]), "key": key, "label": label})
    for j, cell in enumerate(header_cells):
        if _is_blank(cell):
            continue
        key = mapping.get(j)
        if key == "row_no":
            continue
        if key is None or key in _IGNORED_KEYS or (key in key_cols and key_cols[key] != j):
            sheet.ignored.append(_text(cell))
    sheet.has_group_column = "group_name" in key_cols

    for i in range(header_idx + 1, len(grid)):
        row = grid[i]
        if all(_is_blank(c) for c in row):
            continue
        cells = {
            key: row[j] if j < len(row) else None
            for key, j in key_cols.items()
            if key not in _IGNORED_KEYS
        }
        # A "total"/signature line under the table has text only in some far
        # column — require at least a name part or an ИИН to call it a student.
        if not any(
            not _is_blank(cells.get(k)) for k in ("full_name", "last_name", "first_name", "iin")
        ):
            continue
        sheet.rows.append((i + 1, cells))
    return sheet


_PSEUDO_LABELS = {
    "last_name": "Фамилия", "first_name": "Имя", "middle_name": "Отчество",
    "partner_bin": "БИН предприятия",
}


# ─────────────────────────────────────────────────────────────────────────────
# Plan
# ─────────────────────────────────────────────────────────────────────────────

@dataclass
class ImportOptions:
    mode: str = "fill"
    default_group: str = ""
    force_group: str = ""
    group_overrides: dict[str, str] = field(default_factory=dict)  # sheet name → group
    defaults: dict[str, Any] = field(default_factory=dict)         # attr → typed value
    # {sheet name or "*": {header text: field key, "" = do not load}}
    column_overrides: dict[str, dict[str, str]] = field(default_factory=dict)

    @classmethod
    def from_request(cls, raw: dict) -> "ImportOptions":
        mode = (raw.get("mode") or "fill").strip()
        if mode not in MODES:
            raise ImportFileError(f"Неизвестный режим импорта: {mode}")
        opts = cls(
            mode=mode,
            default_group=normalize_group(raw.get("default_group") or ""),
            force_group=normalize_group(raw.get("force_group") or ""),
            group_overrides={
                str(k): normalize_group(v or "")
                for k, v in (raw.get("group_overrides") or {}).items()
                if (v or "").strip()
            },
        )
        allowed = set(FIELD_BY_KEY) | set(_PSEUDO_LABELS) | {""}
        for sheet_name, cols in (raw.get("column_overrides") or {}).items():
            if not isinstance(cols, dict):
                continue
            for header, key in cols.items():
                if key not in allowed:
                    raise ImportFileError(f"Неизвестное поле для колонки «{header}»: {key}")
            opts.column_overrides[str(sheet_name)] = {str(h): str(k) for h, k in cols.items()}
        d = raw.get("defaults") or {}
        for key in DEFAULT_KEYS:
            if _is_blank(d.get(key)):
                continue
            v = d[key]
            try:
                if key in ("practice_start", "practice_end"):
                    opts.defaults[key] = _parse_date(v)
                elif key in ("course", "enrollment_year", "partner_id"):
                    opts.defaults[key] = _parse_int(v)
                else:
                    opts.defaults[key] = _text(v)[: FIELD_BY_KEY[key].max_len or 500]
            except ValueError as exc:
                raise ImportFileError(f"Некорректное значение «{key}» в настройках загрузки: {exc}") from exc
        ds, de = opts.defaults.get("practice_start"), opts.defaults.get("practice_end")
        if ds and de and de < ds:
            raise ImportFileError("Окончание практики не может быть раньше начала")
        pid = opts.defaults.get("partner_id")
        if pid is not None and not db.session.get(Partner, pid):
            raise ImportFileError(f"Партнёр с id={pid} не найден")
        return opts


def available_fields() -> list[dict]:
    """Fields a spreadsheet column can be mapped to (for the manual mapping UI)."""
    out = [{"key": f.key, "label": f.label} for f in FIELDS]
    out += [{"key": k, "label": v} for k, v in _PSEUDO_LABELS.items()]
    return out


def _empty(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _jsonable(v: Any) -> Any:
    if isinstance(v, (date, datetime)):
        return v.isoformat()[:10]
    return v


@dataclass
class PlanRow:
    row: int
    sheet: str
    status: str = "new"            # new | update | unchanged | skip | error
    full_name: str = ""
    iin: str = ""
    group_name: str = ""
    values: dict[str, Any] = field(default_factory=dict)   # Student attr → typed value
    changes: list[dict] = field(default_factory=list)      # update only
    student_id: int | None = None
    match: str | None = None       # iin | name
    derived: set[str] = field(default_factory=set)         # computed, never overwrite
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def to_dict(self, partner_names: dict[int, str]) -> dict:
        v = dict(self.values)
        pid = v.pop("partner_id", None)
        return {
            "row": self.row,
            "sheet": self.sheet,
            "status": self.status,
            "full_name": self.full_name,
            "iin": self.iin,
            "group_name": self.group_name,
            "specialty": v.get("specialty") or "",
            "course": v.get("course"),
            "phone": v.get("phone") or "",
            "partner_name": partner_names.get(pid, "") if pid else "",
            "practice_start": _jsonable(v.get("practice_start")),
            "practice_end": _jsonable(v.get("practice_end")),
            "is_grant_student": bool(v.get("is_grant_student")),
            "student_id": self.student_id,
            "match": self.match,
            "changes": self.changes,
            "errors": self.errors,
            "warnings": self.warnings,
        }


@dataclass
class Plan:
    rows: list[PlanRow]
    sheets: list[dict]
    mode: str
    filename: str
    partner_names: dict[int, str]

    @property
    def counts(self) -> dict[str, int]:
        out = {"total": len(self.rows), "new": 0, "update": 0, "unchanged": 0, "skip": 0, "error": 0}
        for r in self.rows:
            out[r.status] += 1
        return out

    def summary(self) -> dict:
        counts = self.counts
        groups: dict[str, dict] = {}
        for r in self.rows:
            g = groups.setdefault(
                r.group_name or "",
                {"name": r.group_name or "", "total": 0, "new": 0, "update": 0,
                 "unchanged": 0, "skip": 0, "error": 0},
            )
            g["total"] += 1
            g[r.status] += 1
        existing = {
            group_key(name): n
            for name, n in db.session.query(Student.group_name, db.func.count(Student.id))
            .filter(Student.group_name.isnot(None)).group_by(Student.group_name).all()
        }
        group_list = []
        for g in sorted(groups.values(), key=lambda x: (x["name"] == "", x["name"])):
            g["in_db"] = existing.get(group_key(g["name"]), 0) if g["name"] else 0
            group_list.append(g)
        return {
            **counts,
            "with_warnings": sum(1 for r in self.rows if r.warnings and r.status != "error"),
            "no_group": sum(1 for r in self.rows if not r.group_name and r.status != "error"),
            "no_iin": sum(1 for r in self.rows if not r.iin and r.status != "error"),
            "importable": counts["new"] + counts["update"],
            "groups": group_list,
        }

    def to_dict(self) -> dict:
        return {
            "filename": self.filename,
            "mode": self.mode,
            "available_fields": available_fields(),
            "sheets": self.sheets,
            "summary": self.summary(),
            "rows": [r.to_dict(self.partner_names) for r in self.rows],
        }


def _find_partner(raw: str, bin_raw: str, by_bin: dict[str, Partner], by_name: dict[str, list[Partner]],
                  all_partners: list[Partner]) -> tuple[Partner | None, str | None]:
    """Resolve a partner from a name and/or БИН. Returns (partner, problem)."""
    bin_digits = normalize_iin(bin_raw) or (normalize_iin(raw) if len(normalize_iin(raw)) == 12 else "")
    if bin_digits and bin_digits in by_bin:
        return by_bin[bin_digits], None
    key = normalize_name(raw)
    if not key:
        return None, f"БИН {bin_digits} не найден в реестре партнёров" if bin_digits else None
    exact = by_name.get(key) or []
    if len(exact) == 1:
        return exact[0], None
    if len(exact) > 1:
        return None, f"Предприятие «{raw}» найдено несколько раз в реестре — укажите БИН"
    fuzzy = [p for p in all_partners if key in normalize_name(p.organization_name)
             or normalize_name(p.organization_name) in key]
    if len(fuzzy) == 1:
        return fuzzy[0], None
    if len(fuzzy) > 1:
        return None, f"Предприятие «{raw}» подходит под несколько записей реестра — укажите БИН"
    return None, f"Предприятие «{raw}» не найдено в реестре «Партнёры» — поле пропущено"


def build_plan(book: ParsedBook, opts: ImportOptions) -> Plan:
    today = utc_today()
    students = Student.query.all()
    by_iin: dict[str, Student] = {s.iin: s for s in students if s.iin}
    by_name: dict[str, list[Student]] = {}
    group_spelling: dict[str, dict[str, int]] = {}
    for s in students:
        by_name.setdefault(normalize_name(s.full_name), []).append(s)
        if s.group_name:
            spell = group_spelling.setdefault(group_key(s.group_name), {})
            spell[s.group_name] = spell.get(s.group_name, 0) + 1
    canonical_group = {k: max(v.items(), key=lambda kv: kv[1])[0] for k, v in group_spelling.items()}

    all_partners = Partner.query.all()
    partner_by_bin = {p.bin: p for p in all_partners if p.bin}
    partner_by_name: dict[str, list[Partner]] = {}
    for p in all_partners:
        partner_by_name.setdefault(normalize_name(p.organization_name), []).append(p)
    partner_names = {p.id: p.organization_name for p in all_partners}

    specialty_codes: dict[str, str] = {}
    for sp in Specialty.query.filter_by(is_active=True).all():
        if sp.code:
            specialty_codes.setdefault(normalize_name(sp.name), sp.code)

    file_group: dict[str, str] = {}
    seen_iin: dict[str, str] = {}
    seen_name_group: dict[tuple[str, str], str] = {}
    rows: list[PlanRow] = []
    sheet_info: list[dict] = []

    for sheet in book.sheets:
        info = {
            "name": sheet.name,
            "header_row": sheet.header_row,
            "rows": len(sheet.rows),
            "has_group_column": sheet.has_group_column,
            "title_group": sheet.title_group,
            "headerless": sheet.headerless,
            "meta": {k: _jsonable(v) for k, v in sheet.meta.items()},
            "suggested_group": "",
            "group_source": "none",
            "columns": sheet.columns,
            "ignored_columns": sheet.ignored,
            "skipped_reason": sheet.skipped_reason,
        }
        override = opts.group_overrides.get(sheet.name, "")
        if opts.force_group:
            info["group_source"], info["suggested_group"] = "forced", opts.force_group
        elif sheet.has_group_column:
            info["group_source"] = "column"
        elif override:
            info["group_source"], info["suggested_group"] = "override", override
        elif sheet.title_group:
            info["group_source"], info["suggested_group"] = "title", sheet.title_group
        elif not _GENERIC_SHEET.match(sheet.name.strip()):
            info["group_source"], info["suggested_group"] = "sheet", normalize_group(sheet.name)
        elif opts.default_group:
            info["group_source"], info["suggested_group"] = "default", opts.default_group
        sheet_info.append(info)

        sheet_fallback = info["suggested_group"] if info["group_source"] != "column" else ""

        for row_no, cells in sheet.rows:
            pr = PlanRow(row=row_no, sheet=sheet.name)
            rows.append(pr)
            vals: dict[str, Any] = {}

            # ── ФИО ─────────────────────────────────────────────────────────
            name_raw = _text(cells.get("full_name"))
            if not name_raw:
                name_raw = " ".join(
                    p for p in (_text(cells.get("last_name")), _text(cells.get("first_name")),
                                _text(cells.get("middle_name"))) if p
                )
            name = clean_full_name(name_raw)
            pr.full_name = name
            if not name:
                pr.errors.append("Не указано ФИО")
            elif re.search(r"\d", name):
                pr.errors.append(f"В ФИО есть цифры: «{name}» — вероятно, сдвинуты колонки")
            elif len(name) > 200:
                pr.errors.append("ФИО длиннее 200 символов")
            elif " " not in name:
                pr.warnings.append("В ФИО одно слово — проверьте, что указаны фамилия и имя")
            vals["full_name"] = name

            # ── ИИН ─────────────────────────────────────────────────────────
            iin = normalize_iin(cells.get("iin"))
            if iin:
                if len(iin) == 11:
                    iin = "0" + iin
                    pr.warnings.append("В ИИН не хватало цифры — Excel отбросил ведущий 0, он восстановлен")
                if len(iin) != 12:
                    pr.errors.append(f"ИИН должен состоять из 12 цифр (в файле {len(iin)})")
                elif not iin_checksum_ok(iin):
                    pr.warnings.append("Контрольная цифра ИИН не сходится — проверьте номер")
            pr.iin = iin
            where = f"{sheet.name}, строка {row_no}"
            if len(iin) == 12:
                if iin in seen_iin:
                    pr.errors.append(f"Дубль ИИН внутри файла ({seen_iin[iin]})")
                vals["iin"] = iin
            # A missing ИИН is common (group lists often come without it) and is
            # reported once in ``summary.no_iin`` rather than on every row.

            # ── Группа ──────────────────────────────────────────────────────
            if opts.force_group:
                grp = opts.force_group
            else:
                grp = normalize_group(_text(cells.get("group_name"))) or sheet_fallback
            if grp:
                gk = group_key(grp)
                grp = canonical_group.get(gk) or file_group.setdefault(gk, grp)
                if len(grp) > 60:
                    pr.errors.append("Название группы длиннее 60 символов")
                    grp = grp[:60]
            else:
                pr.warnings.append("Не указана группа")
            pr.group_name = grp
            if grp:
                vals["group_name"] = grp

            # ── Простые поля ────────────────────────────────────────────────
            for key in ("specialty", "specialty_code", "phone", "college_supervisor", "partner_supervisor",
                        "practice_type", "form_of_study", "education_program", "id_card_number",
                        "id_card_issued_by", "home_address", "legal_rep_full_name", "legal_rep_phone", "notes"):
                if key in cells and not _is_blank(cells[key]):
                    fd = FIELD_BY_KEY[key]
                    text = _text(cells[key])
                    if fd.max_len and len(text) > fd.max_len:
                        pr.warnings.append(f"«{fd.label}» обрезано до {fd.max_len} символов")
                        text = text[: fd.max_len]
                    vals[key] = text

            if "legal_rep_iin" in cells and not _is_blank(cells["legal_rep_iin"]):
                rep_iin = normalize_iin(cells["legal_rep_iin"])
                if len(rep_iin) == 11:
                    rep_iin = "0" + rep_iin
                if len(rep_iin) == 12:
                    vals["legal_rep_iin"] = rep_iin
                else:
                    pr.warnings.append("ИИН законного представителя некорректен — поле пропущено")

            for key, label in (("course", "Курс"), ("enrollment_year", "Год поступления")):
                if key in cells and not _is_blank(cells[key]):
                    try:
                        n = _parse_int(cells[key])
                    except ValueError:
                        pr.warnings.append(f"«{label}»: не удалось распознать «{_text(cells[key])}» — пропущено")
                        continue
                    if key == "course" and not 1 <= n <= 6:
                        pr.warnings.append(f"«Курс»: значение {n} вне диапазона 1–6 — пропущено")
                        continue
                    if key == "enrollment_year" and not 1990 <= n <= today.year + 1:
                        pr.warnings.append(f"«Год поступления»: {n} выглядит неверно — пропущено")
                        continue
                    vals[key] = n

            for key in ("birth_date", "practice_start", "practice_end"):
                if key in cells and not _is_blank(cells[key]):
                    try:
                        d = _parse_date(cells[key])
                    except ValueError:
                        pr.warnings.append(
                            f"«{FIELD_BY_KEY[key].label}»: дата «{_text(cells[key])}» не распознана — пропущено"
                        )
                        continue
                    if key == "birth_date" and not (1940 <= d.year <= today.year):
                        pr.warnings.append(f"«Дата рождения»: {d:%d.%m.%Y} выглядит неверно — пропущено")
                        continue
                    vals[key] = d

            if "is_grant_student" in cells and not _is_blank(cells["is_grant_student"]):
                try:
                    flag = _parse_bool(cells["is_grant_student"])
                except ValueError:
                    pr.warnings.append(f"«Грант»: не удалось распознать «{_text(cells['is_grant_student'])}»")
                    flag = None
                if flag:
                    vals["is_grant_student"] = True

            # ── Предприятие ─────────────────────────────────────────────────
            partner_raw = _text(cells.get("partner"))
            bin_raw = _text(cells.get("partner_bin"))
            if partner_raw or bin_raw:
                partner, problem = _find_partner(partner_raw, bin_raw, partner_by_bin, partner_by_name, all_partners)
                if partner:
                    vals["partner_id"] = partner.id
                elif problem:
                    pr.warnings.append(problem)

            # ── Справочник специальностей → код ─────────────────────────────
            if vals.get("specialty") and not vals.get("specialty_code"):
                code = specialty_codes.get(normalize_name(vals["specialty"]))
                if code:
                    vals["specialty_code"] = code
                    pr.derived.add("specialty_code")

            # ── Дата рождения из ИИН ────────────────────────────────────────
            if "birth_date" not in vals and len(iin) == 12:
                born, inferred = birth_date_from_iin(iin, today)
                if born:
                    vals["birth_date"] = born
                    pr.derived.add("birth_date")
                    if inferred:
                        pr.warnings.append("Век рождения определён по возрасту — проверьте дату")

            # ── Блок «Специальность / Курс» над таблицей листа ────────────────
            for key in ("specialty", "specialty_code", "course"):
                if key not in vals and sheet.meta.get(key) not in (None, ""):
                    vals[key] = sheet.meta[key]
                    pr.derived.add(key)

            # ── Значения «для всей загрузки» (приоритет ниже строки файла) ───
            for key, dv in opts.defaults.items():
                if key not in vals and dv is not None and dv != "":
                    vals[key] = dv

            if vals.get("practice_start") and vals.get("practice_end") and vals["practice_end"] < vals["practice_start"]:
                pr.errors.append("Окончание практики раньше начала")

            pr.values = {k: v for k, v in vals.items() if not _empty(v)}

            if pr.errors:
                pr.status = "error"
                continue
            # Only a row that survives validation claims its ИИН — a rejected
            # row must not turn a later, valid row into a "duplicate".
            if len(iin) == 12:
                seen_iin[iin] = where

            # ── Поиск существующего студента ────────────────────────────────
            match: Student | None = None
            if len(iin) == 12 and iin in by_iin:
                match, pr.match = by_iin[iin], "iin"
                if normalize_name(match.full_name) != normalize_name(name):
                    pr.warnings.append(f"В базе под этим ИИН записан «{match.full_name}»")
            else:
                cands = by_name.get(normalize_name(name)) or []
                if len(iin) == 12:
                    # An ИИН we don't know yet: only a card WITHOUT an ИИН is the
                    # same person; a different ИИН means a namesake → new card.
                    same = [c for c in cands if not c.iin]
                    if len(same) > 1 and grp:
                        same = [c for c in same if group_key(c.group_name or "") == group_key(grp)] or same
                    if len(same) == 1:
                        match, pr.match = same[0], "name"
                    elif len(same) > 1:
                        pr.errors.append(f"Несколько карточек «{name}» без ИИН (id: {', '.join(str(c.id) for c in same)}) — сопоставить нельзя")
                    elif cands:
                        pr.warnings.append("В базе есть студент с таким же ФИО, но другим ИИН — создана новая карточка")
                else:
                    if grp and len(cands) > 1:
                        narrowed = [c for c in cands if group_key(c.group_name or "") == group_key(grp)]
                        cands = narrowed or cands
                    if len(cands) == 1:
                        match, pr.match = cands[0], "name"
                    elif len(cands) > 1:
                        pr.errors.append(f"В базе {len(cands)} студентов «{name}» — укажите ИИН, чтобы сопоставить")
                    dup_key = (normalize_name(name), group_key(grp))
                    if dup_key in seen_name_group and not match:
                        pr.errors.append(f"Дубль ФИО в группе внутри файла ({seen_name_group[dup_key]})")
                    else:
                        seen_name_group.setdefault(dup_key, where)

            if pr.errors:
                pr.status = "error"
                continue

            if match is None:
                pr.status = "new"
                continue

            pr.student_id = match.id
            if opts.mode == "skip":
                pr.status = "skip"
                continue
            pr.changes = _diff(match, pr.values, opts.mode, pr)
            pr.status = "update" if pr.changes else "unchanged"

    return Plan(rows=rows, sheets=sheet_info, mode=opts.mode, filename=book.filename, partner_names=partner_names)


def _label(attr: str) -> str:
    for f in FIELDS:
        if f.student_attr == attr:
            return f.label
    return attr


def _diff(student: Student, values: dict[str, Any], mode: str, pr: PlanRow) -> list[dict]:
    """What would change on an existing card.

    ``fill``      — only empty fields are written; a differing group is reported.
    ``overwrite`` — non-empty file values replace DB values, except computed
                    values (``pr.derived``) and an already-set ИИН (identity key).
    """
    changes: list[dict] = []
    for attr, new in values.items():
        old = getattr(student, attr, None)
        if attr == "is_grant_student":
            if new and not old:
                changes.append({"field": attr, "label": _label(attr), "old": False, "new": True})
            continue
        if _empty(new) or old == new:
            continue
        keep_existing = (
            mode == "fill"
            or attr in pr.derived
            or attr == "iin"
        )
        if keep_existing and not _empty(old):
            if attr == "group_name":
                pr.warnings.append(
                    f"В базе группа «{old}», в файле «{new}» — не изменена (режим «дополнить пустые»)"
                )
            continue
        changes.append({
            "field": attr, "label": _label(attr),
            "old": _jsonable(old), "new": _jsonable(new),
        })
    return changes


# ─────────────────────────────────────────────────────────────────────────────
# Apply
# ─────────────────────────────────────────────────────────────────────────────

_ASSIGNABLE = {f.student_attr for f in FIELDS} | {"partner_id"}


def apply_plan(plan: Plan, *, year: int | None = None) -> dict:
    """Write the plan inside the caller's transaction (caller commits)."""
    year = year or utc_today().year
    created = updated = 0
    touched_groups: set[str] = set()
    students_by_id: dict[int, Student] = {}
    need = [r.student_id for r in plan.rows if r.status == "update" and r.student_id]
    if need:
        students_by_id = {s.id: s for s in Student.query.filter(Student.id.in_(need)).all()}

    for r in plan.rows:
        if r.status == "new":
            data = {k: v for k, v in r.values.items() if k in _ASSIGNABLE}
            data.setdefault("form_of_study", "очная")
            data.setdefault("practice_type", "профессиональной")
            if "enrollment_year" not in data and data.get("course"):
                data["enrollment_year"] = year - int(data["course"]) + 1
            db.session.add(Student(**data))
            created += 1
        elif r.status == "update":
            s = students_by_id.get(r.student_id)
            if s is None:
                continue
            for ch in r.changes:
                setattr(s, ch["field"], r.values[ch["field"]])
            updated += 1
        else:
            continue
        if r.group_name:
            touched_groups.add(r.group_name)
    return {"created": created, "updated": updated, "groups": sorted(touched_groups)}


# ─────────────────────────────────────────────────────────────────────────────
# Template / export workbooks
# ─────────────────────────────────────────────────────────────────────────────

_TEXT_ATTRS = {"iin", "phone", "legal_rep_iin", "legal_rep_phone", "id_card_number"}
_TEMPLATE_ROWS = 600  # pre-formatted blank rows


def _column_index(key: str) -> int:
    return [f.key for f in FIELDS].index(key) + 1


def build_template_workbook(*, partners: list[Partner], specialties: list[Specialty],
                            groups: list[str]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Студенты"
    headers = [f.label for f in FIELDS]
    ws.append(headers)
    style_header(ws, 1, len(headers), required=[f.required for f in FIELDS])
    for idx, f in enumerate(FIELDS, start=1):
        if f.hint:
            ws.cell(row=1, column=idx).comment = Comment(f.hint, "CCU PRACTICUM", width=260, height=90)
        letter = ws.cell(row=1, column=idx).column_letter
        ws.column_dimensions[letter].width = max(14, min(34, len(f.label) + 4))
        fmt = "@" if f.student_attr in _TEXT_ATTRS else ("DD.MM.YYYY" if f.kind == "date" else None)
        if fmt:
            for r in range(2, _TEMPLATE_ROWS + 2):
                ws.cell(row=r, column=idx).number_format = fmt
    ws.freeze_panes = "D2"
    ws.row_dimensions[1].height = 46

    # ── справочники для выпадающих списков ──────────────────────────────────
    ref = wb.create_sheet("Справочники")
    ref.append(["Предприятия (из реестра)", "Специальности (из «Данные»)", "Группы (уже в системе)"])
    style_header(ref, 1, 3, required=[False, False, False])
    for i in range(max(len(partners), len(specialties), len(groups))):
        ref.append([
            partners[i].organization_name if i < len(partners) else None,
            specialties[i].name if i < len(specialties) else None,
            groups[i] if i < len(groups) else None,
        ])
    autosize(ref, max_width=60)

    last = _TEMPLATE_ROWS + 1

    def _list_validation(key: str, formula: str, *, strict: bool, title: str) -> None:
        col = ws.cell(row=1, column=_column_index(key)).column_letter
        dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
        dv.errorTitle, dv.error = title, "Выберите значение из списка"
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{last}")

    if partners:
        _list_validation("partner", f"'Справочники'!$A$2:$A${len(partners) + 1}", strict=False, title="Предприятие")
    if specialties:
        _list_validation("specialty", f"'Справочники'!$B$2:$B${len(specialties) + 1}", strict=False, title="Специальность")
    if groups:
        _list_validation("group_name", f"'Справочники'!$C$2:$C${len(groups) + 1}", strict=False, title="Группа")
    _list_validation("is_grant_student", '"Да,Нет"', strict=True, title="Грант")
    _list_validation("form_of_study", '"очная,заочная,вечерняя"', strict=False, title="Форма обучения")
    _list_validation("practice_type", '"учебной,профессиональной,производственной,преддипломной"',
                     strict=False, title="Вид практики")
    course_col = ws.cell(row=1, column=_column_index("course")).column_letter
    cdv = DataValidation(type="whole", operator="between", formula1="1", formula2="6",
                         allow_blank=True, showErrorMessage=True)
    cdv.errorTitle, cdv.error = "Курс", "Курс — число от 1 до 6"
    ws.add_data_validation(cdv)
    cdv.add(f"{course_col}2:{course_col}{last}")

    # ── инструкция ──────────────────────────────────────────────────────────
    info = wb.create_sheet("Инструкция", 0)
    info.sheet_view.showGridLines = False
    lines = [
        ("Загрузка студентов по группам", "title"),
        ("", None),
        ("Как заполнить", "h"),
        ("1. На листе «Студенты» одна строка = один студент. Обязательные колонки — коралловые: ФИО, ИИН, Группа.", None),
        ("2. Группа: впишите в колонку «Группа» либо создайте отдельный лист на каждую группу и назовите лист группой (например «ПО 25-1»).", None),
        ("3. Колонки «Предприятие», «Начало/Окончание практики» и руководителей можно не заполнять — их можно задать сразу для всей группы в мастере загрузки.", None),
        ("4. ИИН и телефоны храните как текст (в шаблоне это уже настроено), иначе Excel отбросит ведущий ноль.", None),
        ("5. Даты — в формате ДД.ММ.ГГГГ. Если «Дата рождения» пуста, она берётся из ИИН.", None),
        ("6. Не удаляйте и не переименовывайте заголовки. Лишние колонки система проигнорирует и покажет вам.", None),
        ("", None),
        ("Что происходит при загрузке", "h"),
        ("• Сначала система показывает проверку: сколько студентов будет добавлено, обновлено и какие строки с ошибками. Ничего не сохраняется, пока вы не подтвердите.", None),
        ("• Студент ищется в базе по ИИН, затем по ФИО. Дубли не создаются.", None),
        ("• Режимы для уже существующих студентов: «только новые», «дополнить пустые поля» (по умолчанию, ничего не затирает) и «обновить данными из файла».", None),
        ("• Предприятие должно быть заранее добавлено в раздел «Партнёры» — выберите его из выпадающего списка (БИН тоже подойдёт).", None),
    ]
    for r, (text, style) in enumerate(lines, start=1):
        cell = info.cell(row=r, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if style == "title":
            cell.font = Font(bold=True, size=16, color="E85A3F")
        elif style == "h":
            cell.font = Font(bold=True, size=12, color="46484B")
    info.column_dimensions["A"].width = 120
    wb.active = 1
    return wb


def build_students_workbook(students: list[Student], contract_info: dict[int, dict] | None = None,
                            partner_names: dict[int, str] | None = None) -> Workbook:
    """Export in the SAME column order as the template, so an exported group
    can be edited in Excel and loaded back; the trailing two columns are
    informational and ignored by the importer."""
    contract_info = contract_info or {}
    partner_names = partner_names or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Студенты"
    headers = [f.label for f in FIELDS] + ["Договор практики №", "Статус договора"]
    rows = []
    for s in students:
        row = []
        for f in FIELDS:
            if f.key == "partner":
                row.append(partner_names.get(s.partner_id, "") if s.partner_id else "")
                continue
            v = getattr(s, f.student_attr, None)
            if f.kind == "bool":
                v = "Да" if v else ""
            elif isinstance(v, date):
                v = v.strftime("%d.%m.%Y")
            row.append(v)
        ci = contract_info.get(s.id) or {}
        row.extend([ci.get("number", ""), ci.get("status_label", "")])
        rows.append(row)
    text_cols = [i for i, f in enumerate(FIELDS, start=1) if f.student_attr in _TEXT_ATTRS]
    table_sheet(
        ws, headers, rows,
        required=[f.required for f in FIELDS] + [False, False],
        text_columns=text_cols,
    )
    return wb
