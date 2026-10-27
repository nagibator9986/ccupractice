"""Small helpers for producing styled .xlsx downloads (registers, link lists,
the student import template) so every export looks the same."""
from __future__ import annotations

import io
from typing import Iterable, Sequence
from urllib.parse import quote

from flask import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

CORAL = "E85A3F"
CHARCOAL = "46484B"
_HEADER_FILL = PatternFill("solid", fgColor=CORAL)
_OPTIONAL_FILL = PatternFill("solid", fgColor=CHARCOAL)
_THIN = Side(style="thin", color="D9DADD")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def style_header(ws, row: int, ncols: int, *, required: Sequence[bool] | None = None) -> None:
    """Bold white header cells; coral for required columns, charcoal otherwise."""
    for idx in range(1, ncols + 1):
        cell = ws.cell(row=row, column=idx)
        is_req = True if required is None else bool(required[idx - 1])
        cell.fill = _HEADER_FILL if is_req else _OPTIONAL_FILL
        cell.font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER
    ws.row_dimensions[row].height = 32


def autosize(ws, *, min_width: int = 10, max_width: int = 48, header_row: int = 1) -> None:
    """Width from the longest rendered value per column (capped)."""
    widths: dict[int, int] = {}
    for row in ws.iter_rows(min_row=header_row, max_row=min(ws.max_row, 400)):
        for cell in row:
            if cell.value is None:
                continue
            longest = max(len(part) for part in str(cell.value).split("\n"))
            widths[cell.column] = max(widths.get(cell.column, 0), longest)
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_width, min(max_width, width + 3))


def table_sheet(
    ws,
    headers: Sequence[str],
    rows: Iterable[Sequence],
    *,
    required: Sequence[bool] | None = None,
    freeze: bool = True,
    text_columns: Sequence[int] = (),
) -> int:
    """Write a header + rows table. ``text_columns`` are 1-based indexes forced
    to text format (ИИН / телефоны: Excel would otherwise drop leading zeros).
    Returns the number of data rows written."""
    ws.append(list(headers))
    style_header(ws, 1, len(headers), required=required)
    count = 0
    for r in rows:
        ws.append(list(r))
        count += 1
    text_cols = set(text_columns)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.border = _BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=False)
            if cell.column in text_cols and cell.value is not None:
                cell.number_format = "@"
    if freeze:
        ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(ws.max_row, 1)}"
    autosize(ws)
    return count


def workbook_response(wb: Workbook, filename: str) -> Response:
    """Serialise to an attachment response (RFC 5987 filename for Cyrillic)."""
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    ascii_name = "".join(ch if ord(ch) < 128 else "_" for ch in filename)
    resp = Response(buf.getvalue(), mimetype=XLSX_MIME)
    resp.headers["Content-Disposition"] = (
        f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(filename)}"
    )
    resp.headers["Cache-Control"] = "no-store"
    resp.headers["Access-Control-Expose-Headers"] = "Content-Disposition"
    return resp
