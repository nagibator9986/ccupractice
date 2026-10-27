from __future__ import annotations

from datetime import date, timedelta
from io import BytesIO

from openpyxl import load_workbook

from tests.base import AppTestCase


class PracticeBulkTests(AppTestCase):
    def setUp(self):
        self.reset_db()
        from app.extensions import db
        from app.models import Partner, Student, SigningRequest

        with self.app.app_context():
            SigningRequest.query.delete()
            self.partner = Partner(
                organization_name="ТОО Ромашка", bin="123456789012", legal_address="г. Алматы",
                director_full_name="Директоров Д.Д.", seats_count=2, phone="+7 701 000 00 00",
            )
            db.session.add(self.partner)
            db.session.flush()
            self.pid = self.partner.id
            ids = []
            for i in range(3):
                s = Student(
                    full_name=f"Студент{i} Тест", iin=self.make_iin("20080101", "5", f"{i + 1:04d}"),
                    group_name="ПО 25-1", practice_start=date(2026, 6, 1), practice_end=date(2026, 6, 30),
                    phone=f"+7 777 000 00 0{i}",
                )
                db.session.add(s)
                db.session.flush()
                ids.append(s.id)
            # one student with gaps (no ИИН / dates)
            incomplete = Student(full_name="Пустой Студент", group_name="ПО 25-1")
            db.session.add(incomplete)
            db.session.flush()
            self.incomplete_id = incomplete.id
            self.ids = ids
            db.session.commit()

    def post(self, path, body=None, headers=None):
        return self.client.post(path, json=body or {}, headers=headers or self.admin)

    # ── bulk contract creation ─────────────────────────────────────────────
    def test_bulk_dry_run_then_create_skips_incomplete_and_existing(self):
        body = {"group": "ПО 25-1", "partner_id": self.pid, "contract_date": "2026-05-20"}
        dry = self.post("/api/contracts/bulk", {**body, "dry_run": True}).get_json()
        self.assertEqual((dry["will_create"], dry["skipped"], dry["created"]), (3, 1, 0))
        skipped = [i for i in dry["items"] if i["action"] == "skip"][0]
        self.assertEqual(skipped["code"], "incomplete")
        self.assertIn("ИИН студента", skipped["reason"])
        # 3 contracts > 2 seats → capacity warning
        self.assertTrue(any(w["code"] == "partner_overbooked" for w in dry["warnings"]))

        r = self.post("/api/contracts/bulk", body)
        self.assertEqual(r.status_code, 201, r.get_json())
        res = r.get_json()
        self.assertEqual(res["created"], 3)
        numbers = [i["number"] for i in res["items"] if i["number"]]
        self.assertEqual(numbers, sorted(numbers))                   # sequential numbering
        self.assertEqual(len(set(numbers)), 3)

        again = self.post("/api/contracts/bulk", body).get_json()   # idempotent
        self.assertEqual((again["created"], again["skipped"]), (0, 4))
        self.assertEqual({i["code"] for i in again["items"]}, {"has_contract", "incomplete"})

    def test_bulk_uses_students_own_partner_and_reports_missing_partner(self):
        from app.extensions import db
        from app.models import Student

        with self.app.app_context():
            s = db.session.get(Student, self.ids[0])
            s.partner_id = self.pid
            db.session.commit()
        res = self.post("/api/contracts/bulk", {"student_ids": self.ids}).get_json()
        by_id = {i["student_id"]: i for i in res["items"]}
        self.assertEqual(by_id[self.ids[0]]["action"], "create")
        self.assertEqual(by_id[self.ids[1]]["code"], "no_partner")

    def test_bulk_requires_admin_and_selection(self):
        self.assertEqual(self.post("/api/contracts/bulk", {"group": "ПО 25-1"}, self.viewer).status_code, 403)
        self.assertEqual(self.post("/api/contracts/bulk", {}).status_code, 400)

    # ── single create / duplicates / readiness ─────────────────────────────
    def test_duplicate_guard_and_readiness_gate(self):
        sid = self.ids[0]
        r = self.post("/api/contracts", {"student_id": sid, "partner_id": self.pid, "contract_date": "2026-05-20"})
        self.assertEqual(r.status_code, 201)
        cid = r.get_json()["item"]["id"]
        self.assertTrue(r.get_json()["item"]["readiness"]["ok"])
        dup = self.post("/api/contracts", {"student_id": sid, "partner_id": self.pid, "contract_date": "2026-05-21"})
        self.assertEqual((dup.status_code, dup.get_json()["code"]), (409, "duplicate_contract"))
        ok = self.post("/api/contracts", {"student_id": sid, "partner_id": self.pid, "allow_duplicate": True})
        self.assertEqual(ok.status_code, 201)
        newest = ok.get_json()["item"]["id"]

        # incomplete student: generation is refused with the list of gaps
        r = self.post("/api/contracts", {"student_id": self.incomplete_id, "partner_id": self.pid})
        icid = r.get_json()["item"]["id"]
        gen = self.post(f"/api/contracts/{icid}/generate")
        self.assertEqual((gen.status_code, gen.get_json()["code"]), (422, "incomplete"))
        labels = {m["label"] for m in gen.get_json()["readiness"]["missing"] if m["severity"] == "critical"}
        self.assertIn("ИИН студента", labels)

        pre = self.client.get(f"/api/contracts/readiness?student_id={sid}&partner_id={self.pid}", headers=self.admin).get_json()
        self.assertEqual(pre["duplicate"]["id"], newest)          # newest twin is reported
        self.assertNotEqual(newest, cid)

    def test_next_action_and_delete_guard(self):
        from app.extensions import db
        from app.models import Contract, Signature

        cid = self.post("/api/contracts", {"student_id": self.ids[0], "partner_id": self.pid}).get_json()["item"]["id"]
        self.assertEqual(self.client.get(f"/api/contracts/{cid}", headers=self.admin).get_json()["item"]["next_action"], "generate")
        with self.app.app_context():
            c = db.session.get(Contract, cid)
            c.docx_path = "x.docx"
            db.session.add(Signature(contract_id=cid, signer_role="student", cms_signature="x", signed_payload_sha256="y"))
            db.session.commit()
        item = self.client.get(f"/api/contracts/{cid}", headers=self.admin).get_json()["item"]
        self.assertEqual((item["next_action"], item["signed_roles"]), ("await_signatures", ["student"]))
        r = self.client.delete(f"/api/contracts/{cid}", headers=self.admin)
        self.assertEqual((r.status_code, r.get_json()["code"]), (409, "has_signatures"))
        r = self.client.delete(f"/api/contracts/{cid}", json={"force": True}, headers=self.admin)
        self.assertEqual(r.status_code, 200)

    # ── signing links ──────────────────────────────────────────────────────
    def _contract_with_docx(self, sid):
        from app.extensions import db
        from app.models import Contract
        import os

        cid = self.post("/api/contracts", {"student_id": sid, "partner_id": self.pid}).get_json()["item"]["id"]
        with self.app.app_context():
            c = db.session.get(Contract, cid)
            archive = self.app.config["ARCHIVE_FOLDER"]
            os.makedirs(archive, exist_ok=True)
            path = f"{archive}/c{cid}.docx"
            with open(path, "wb") as fh:
                fh.write(b"docx-bytes")
            c.docx_path = f"c{cid}.docx"
            c.status = "generated"
            db.session.commit()
        return cid

    def test_bulk_invite_is_idempotent_and_exports_links(self):
        cids = [self._contract_with_docx(s) for s in self.ids[:2]]
        no_docx = self.post("/api/contracts", {"student_id": self.ids[2], "partner_id": self.pid}).get_json()["item"]["id"]
        res = self.post("/api/signing/bulk-invite", {"contract_ids": cids + [no_docx]}).get_json()
        self.assertEqual(res["new_links"], 4)                       # partner + student × 2
        statuses = {i["contract_id"]: i["status"] for i in res["items"]}
        self.assertEqual(statuses[no_docx], "skipped")
        link = res["items"][0]["links"][0]
        self.assertTrue(link["is_new"])
        self.assertIn("/sign/", link["sign_url"])
        self.assertTrue(link["recipient_phone"])

        again = self.post("/api/signing/bulk-invite", {"contract_ids": cids}).get_json()
        self.assertEqual(again["new_links"], 0)                     # no second live token
        self.assertFalse(any(l["is_new"] for i in again["items"] for l in i["links"]))

        x = self.post("/api/signing/links-export", {"contract_ids": cids})
        self.assertEqual(x.status_code, 200)
        ws = load_workbook(BytesIO(x.data)).active
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(len(rows), 5)                              # header + 4 links
        self.assertIn("/sign/", rows[1][7])
        self.assertIn("Здравствуйте", rows[1][10])
        self.assertEqual(self.post("/api/signing/bulk-invite", {"contract_ids": cids}, self.viewer).status_code, 403)

    def test_public_view_masks_identifiers_and_extend_keeps_token(self):
        cid = self._contract_with_docx(self.ids[0])
        link = self.post("/api/signing/bulk-invite", {"contract_ids": [cid], "roles": ["student"]}).get_json()["items"][0]["links"][0]
        token = link["token"]
        pub = self.client.get(f"/api/signing/public/{token}").get_json()
        self.assertIn("iin_masked", pub["contract"]["student"])
        self.assertNotIn("iin", pub["contract"]["student"])
        self.assertIn("****", pub["request"]["recipient_iin_or_bin_masked"])
        self.assertIsNone(pub["download"]["pdf_preview"])           # no PDF yet
        self.assertIn("phone", pub["contract"]["college"])
        ext = self.post(f"/api/signing/requests/{link['id']}/extend", {"days": 90}).get_json()["item"]
        self.assertEqual(ext["token"], token)
        self.assertEqual(self.client.get(f"/api/signing/public/{token}").status_code, 200)

    # ── lists / stats / export ─────────────────────────────────────────────
    def test_students_filters_groups_and_contract_summary(self):
        self.post("/api/contracts/bulk", {"group": "ПО 25-1", "partner_id": self.pid})
        none = self.client.get("/api/students?contract=none&with_contracts=1", headers=self.admin).get_json()["items"]
        self.assertEqual([s["full_name"] for s in none], ["Пустой Студент"])
        anyc = self.client.get("/api/students?contract=any&with_contracts=1", headers=self.admin).get_json()["items"]
        self.assertEqual(len(anyc), 3)
        self.assertTrue(all(s["contract"]["number"] for s in anyc))
        groups = self.client.get("/api/students/groups", headers=self.admin).get_json()["items"]
        g = groups[0]
        self.assertEqual((g["name"], g["students"], g["contracts"], g["without_contract"]), ("ПО 25-1", 4, 3, 1))
        self.assertEqual(g["by_status"]["draft"], 3)

    def test_group_filter_folds_spellings_and_handles_none(self):
        from app.extensions import db
        from app.models import Student

        with self.app.app_context():
            db.session.add(Student(full_name="Вариант Написания", group_name="ПО 25 - 1"))
            db.session.add(Student(full_name="Без Группы"))
            db.session.commit()
        got = self.client.get("/api/students?group=ПО 25-1", headers=self.admin).get_json()["items"]
        self.assertEqual(len(got), 5)                                  # 4 fixtures + the variant spelling
        none = self.client.get("/api/students?group=__none__", headers=self.admin).get_json()["items"]
        self.assertEqual([s["full_name"] for s in none], ["Без Группы"])
        x = self.client.get("/api/students/export?group=__none__", headers=self.admin)
        rows = list(load_workbook(BytesIO(x.data)).active.iter_rows(values_only=True))
        self.assertEqual((len(rows), rows[1][0]), (2, "Без Группы"))

    def test_bulk_practice_and_delete(self):
        r = self.post("/api/students/bulk-practice", {
            "group": "ПО 25-1", "fields": {"partner_id": self.pid, "practice_start": "2026-07-01",
                                            "practice_end": "2026-07-31", "college_supervisor": "Иванов И.И."}})
        res = r.get_json()
        self.assertEqual((r.status_code, res["updated"]), (200, 4))
        self.assertTrue(any(w["code"] == "partner_overbooked" for w in res["warnings"]))   # 4 > 2 seats
        bad = self.post("/api/students/bulk-practice", {"group": "ПО 25-1", "fields": {"practice_start": "2026-09-01"}}).get_json()
        self.assertEqual(len(bad["skipped"]), 4)                  # end (07-31) < new start (09-01)
        self.assertEqual(self.post("/api/students/bulk-practice", {"group": "ПО 25-1", "fields": {}}).status_code, 400)

        self.post("/api/contracts/bulk", {"student_ids": self.ids[:1], "partner_id": self.pid})
        d = self.post("/api/students/bulk-delete", {"student_ids": self.ids[:2]}).get_json()
        self.assertEqual((d["deleted"], len(d["blocked"])), (1, 1))        # one has a contract
        self.assertEqual(self.post("/api/students/bulk-delete", {"student_ids": [self.ids[2]]}, self.viewer).status_code, 403)

    def test_partner_stats_and_stats_endpoint_and_export(self):
        self.post("/api/contracts/bulk", {"group": "ПО 25-1", "partner_id": self.pid})
        partners = self.client.get("/api/partners?with_stats=1", headers=self.admin).get_json()["items"]
        self.assertEqual((partners[0]["contracts_count"], partners[0]["seats_free"]), (3, 2))
        st = self.client.get("/api/contracts/stats", headers=self.admin).get_json()
        self.assertEqual(st["totals"]["contracts"], 3)
        self.assertEqual(st["by_status"]["draft"], 3)
        self.assertEqual(st["groups"][0]["name"], "ПО 25-1")
        x = self.client.get("/api/contracts/export?group=ПО 25-1", headers=self.admin)
        self.assertEqual(x.status_code, 200)
        self.assertEqual(len(list(load_workbook(BytesIO(x.data)).active.iter_rows())), 4)
        self.assertEqual(self.client.get("/api/contracts/export").status_code, 401)
        multi = self.client.get("/api/contracts?status=draft,signed", headers=self.admin).get_json()
        self.assertEqual(multi["total"], 3)
        self.assertEqual(self.client.get("/api/contracts?status=draft,signed&has_file=1", headers=self.admin).get_json()["total"], 0)


if __name__ == "__main__":
    import unittest

    unittest.main()
