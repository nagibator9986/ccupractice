from __future__ import annotations

from datetime import date

from tests.base import AppTestCase


class StudentImportTests(AppTestCase):
    def setUp(self):
        self.reset_db()

    # ── file → plan ────────────────────────────────────────────────────────
    def test_sheet_per_group_and_group_column(self):
        iin1 = self.make_iin(serial="0001")
        iin2 = self.make_iin("20090101", "5", "0002")
        data = self.xlsx({
            "ПО 25-1": [["№", "ФИО", "ИИН", "Специальность"],
                        [1, "иванов иван иванович", iin1, "Программное обеспечение"]],
            "Список": [["ФИО", "ИИН", "Группа"], ["Петров Пётр", iin2, "ТК 25-2"]],
        })
        r = self.upload("/api/students/import/preview", data)
        self.assertEqual(r.status_code, 200, r.get_json())
        body = r.get_json()
        s = body["summary"]
        self.assertEqual((s["new"], s["error"], s["total"]), (2, 0, 2))
        by_name = {row["full_name"]: row for row in body["rows"]}
        self.assertEqual(by_name["Иванов Иван Иванович"]["group_name"], "ПО 25-1")  # sheet name, caps fixed
        self.assertEqual(by_name["Петров Пётр"]["group_name"], "ТК 25-2")            # group column
        self.assertEqual({g["name"] for g in s["groups"]}, {"ПО 25-1", "ТК 25-2"})

    def test_title_row_names_the_group_and_header_is_found_below(self):
        iin = self.make_iin()
        data = self.xlsx({"Лист1": [["Список студентов группы ПО 24-3"], [],
                                    ["№ п/п", "Ф.И.О.", "ИИН"], [1, "Сидоров Сидор", iin]]})
        body = self.upload("/api/students/import/preview", data).get_json()
        self.assertEqual(body["rows"][0]["group_name"], "ПО 24-3")
        self.assertEqual(body["sheets"][0]["group_source"], "title")

    def test_force_group_and_default_group(self):
        iin = self.make_iin()
        data = self.xlsx({"Лист1": [["ФИО", "ИИН"], ["Сидоров Сидор", iin]]})
        no_group = self.upload("/api/students/import/preview", data).get_json()
        self.assertEqual(no_group["summary"]["no_group"], 1)
        default = self.upload("/api/students/import/preview", data, {"default_group": "ПО 25-1"}).get_json()
        self.assertEqual(default["rows"][0]["group_name"], "ПО 25-1")
        forced = self.upload("/api/students/import/preview", data, {"force_group": "ИС 25-9"}).get_json()
        self.assertEqual(forced["rows"][0]["group_name"], "ИС 25-9")

    def test_excel_artifacts_iin_without_leading_zero_and_float(self):
        iin = self.make_iin("20060812", "5", "0003")           # starts with 0
        self.assertTrue(iin.startswith("0"))
        data = self.xlsx({"ПО 25-1": [["ФИО", "ИИН"], ["Алиев Али", int(iin)], ["Бек Бек", float(self.make_iin("20080101", "5", "0004"))]]})
        body = self.upload("/api/students/import/preview", data).get_json()
        rows = {r["full_name"]: r for r in body["rows"]}
        self.assertEqual(rows["Алиев Али"]["iin"], iin)
        self.assertTrue(any("ведущий 0" in w for w in rows["Алиев Али"]["warnings"]))
        self.assertEqual(rows["Бек Бек"]["status"], "new")

    def test_validation_errors(self):
        iin = self.make_iin()
        data = self.xlsx({"ПО 25-1": [
            ["ФИО", "ИИН", "Начало практики", "Окончание практики"],
            ["", iin, None, None],
            ["Иванов 123", self.make_iin("20080101", "5", "0009"), None, None],
            ["Короткий Иин", "12345", None, None],
            ["Первый Дубль", iin, None, None],
            ["Второй Дубль", iin, None, None],
            ["Даты Наоборот", self.make_iin("20070707", "5", "0011"), "01.06.2026", "01.05.2026"],
        ]})
        rows = self.upload("/api/students/import/preview", data).get_json()["rows"]
        status = {r["row"]: r["status"] for r in rows}
        self.assertEqual(status[2], "error")   # no name
        self.assertEqual(status[3], "error")   # digits in name
        self.assertEqual(status[4], "error")   # short ИИН
        self.assertEqual(status[5], "new")
        self.assertEqual(status[6], "error")   # duplicate ИИН in file
        self.assertEqual(status[7], "error")   # end before start

    def test_bad_files(self):
        r = self.upload("/api/students/import/preview", b"not an excel", filename="x.xlsx")
        self.assertEqual(r.status_code, 400)
        r = self.upload("/api/students/import/preview", b"x", filename="x.xls")
        self.assertIn(".xlsx", r.get_json()["error"])
        r = self.upload("/api/students/import/preview", self.xlsx({"A": [["foo", "bar"], [1, 2]]}))
        self.assertEqual(r.status_code, 400)
        self.assertIn("шаблон", r.get_json()["error"].lower())

    def test_csv_semicolon_cp1251(self):
        iin = self.make_iin()
        csv_bytes = f"ФИО;ИИН;Группа\nКим Анна;{iin};ПО 25-1\n".encode("cp1251")
        body = self.upload("/api/students/import/preview", csv_bytes, filename="g.csv").get_json()
        self.assertEqual(body["summary"]["new"], 1)

    def test_real_college_layout_meta_block_and_headerless_names(self):
        """«Курс / Специальность / Группа / Куратор» block + a bare column of names."""
        data = self.xlsx({"МГК 26-1": [
            ["Курс:", 1], ["Специальность:", "04140100 Маркетинг (по отраслям)"],
            ["Группа:", "МГК 26-1 грант"], ["Куратор:", "Тұрлыбек Жания Ғанибекқыз"],
            [], ["Абдулаева Дария"], ["Амеркулова Жансая Ешмухан"], ["Аметова Рамина Нурахметовна"],
        ]})
        body = self.upload("/api/students/import/preview", data).get_json()
        self.assertEqual(body["summary"]["new"], 3, body)
        sheet = body["sheets"][0]
        self.assertTrue(sheet["headerless"])
        self.assertEqual(sheet["meta"]["curator"], "Тұрлыбек Жания Ғанибекқыз")
        rows = body["rows"]
        self.assertEqual({r["group_name"] for r in rows}, {"МГК 26-1 грант"})
        self.assertEqual({r["specialty"] for r in rows}, {"Маркетинг (по отраслям)"})
        self.assertEqual({r["course"] for r in rows}, {1})
        # specialty code from the block is stored, curator lines are not mistaken for students
        self.assertEqual(self.upload("/api/students/import/commit", data).status_code, 200)
        from app.models import Student
        with self.app.app_context():
            self.assertEqual(Student.query.count(), 3)
            self.assertEqual({s.specialty_code for s in Student.query.all()}, {"04140100"})

    def test_meta_block_above_a_normal_header(self):
        iin = self.make_iin()
        data = self.xlsx({"Лист1": [["Группа", "ПО 26-2"], ["Курс", 2], [],
                                    ["№", "ФИО", "ИИН"], [1, "Сидоров Сидор", iin]]})
        row = self.upload("/api/students/import/preview", data).get_json()["rows"][0]
        self.assertEqual((row["group_name"], row["course"]), ("ПО 26-2", 2))

    def test_manual_column_mapping_for_unrecognised_header(self):
        iin = self.make_iin()
        data = self.xlsx({"ПО 25-1": [["ФИО", "ИИН", "Наставник группы"],
                                      ["Сидоров Сидор", iin, "Иванова А.А."]]})
        first = self.upload("/api/students/import/preview", data).get_json()
        self.assertIn("Наставник группы", first["sheets"][0]["ignored_columns"] + ["Наставник группы"])
        self.assertTrue(any(f["key"] == "college_supervisor" for f in first["available_fields"]))
        mapped = self.upload("/api/students/import/commit", data, {
            "column_overrides": {"ПО 25-1": {"Наставник группы": "college_supervisor"}}})
        self.assertEqual(mapped.status_code, 200, mapped.get_json())
        from app.models import Student
        with self.app.app_context():
            self.assertEqual(Student.query.one().college_supervisor, "Иванова А.А.")
        bad = self.upload("/api/students/import/preview", data, {"column_overrides": {"*": {"X": "nonsense"}}})
        self.assertEqual(bad.status_code, 400)

    # ── commit ─────────────────────────────────────────────────────────────
    def test_commit_creates_and_is_idempotent(self):
        from app.models import Partner, Student
        from app.extensions import db

        with self.app.app_context():
            p = Partner(organization_name="ТОО Ромашка", bin="123456789012")
            db.session.add(p)
            db.session.commit()
            pid = p.id
        iin = self.make_iin()
        data = self.xlsx({"ПО 25-1": [["ФИО", "ИИН", "Курс", "Предприятие"],
                                      ["Иванов Иван", iin, "1 курс", "ромашка"]]})
        opts = {"defaults": {"practice_start": "2026-06-01", "practice_end": "2026-06-30"}}
        r = self.upload("/api/students/import/commit", data, opts, expect_new=1, expect_update=0)
        self.assertEqual(r.status_code, 200, r.get_json())
        self.assertEqual(r.get_json()["created"], 1)
        with self.app.app_context():
            s = Student.query.one()
            self.assertEqual((s.group_name, s.course, s.partner_id), ("ПО 25-1", 1, pid))
            self.assertEqual(s.practice_start, date(2026, 6, 1))
            self.assertEqual(s.birth_date, date(2008, 5, 12))      # derived from ИИН
        again = self.upload("/api/students/import/preview", data, opts).get_json()["summary"]
        self.assertEqual((again["new"], again["unchanged"]), (0, 1))
        r = self.upload("/api/students/import/commit", data, opts)
        self.assertEqual(r.status_code, 422)                       # nothing to do

    def test_modes_fill_vs_overwrite_vs_skip(self):
        from app.models import Student
        from app.extensions import db

        iin = self.make_iin()
        with self.app.app_context():
            db.session.add(Student(full_name="Иванов Иван", iin=iin, group_name="ПО 25-1", phone="111"))
            db.session.commit()
        data = self.xlsx({"ПО 25-1": [["ФИО", "ИИН", "Телефон", "Специальность"],
                                      ["Иванов Иван", iin, "222", "ПО"]]})

        skip = self.upload("/api/students/import/preview", data, {"mode": "skip"}).get_json()
        self.assertEqual(skip["rows"][0]["status"], "skip")

        fill = self.upload("/api/students/import/preview", data, {"mode": "fill"}).get_json()["rows"][0]
        self.assertEqual({c["field"] for c in fill["changes"]}, {"specialty", "birth_date"})   # phone kept

        over = self.upload("/api/students/import/preview", data, {"mode": "overwrite"}).get_json()["rows"][0]
        self.assertIn("phone", {c["field"] for c in over["changes"]})
        self.assertEqual(self.upload("/api/students/import/commit", data, {"mode": "overwrite"}).status_code, 200)
        with self.app.app_context():
            self.assertEqual(Student.query.one().phone, "222")

    def test_name_match_fills_missing_iin_and_namesake_gets_new_card(self):
        from app.models import Student
        from app.extensions import db

        iin_new = self.make_iin("20080303", "5", "0021")
        iin_other = self.make_iin("20080404", "5", "0022")
        with self.app.app_context():
            db.session.add(Student(full_name="Без Иина", group_name="ПО 25-1"))
            db.session.add(Student(full_name="Тёзка Тёзкин", iin=iin_other, group_name="ПО 25-1"))
            db.session.commit()
        data = self.xlsx({"ПО 25-1": [["ФИО", "ИИН"], ["Без Иина", iin_new],
                                      ["Тёзка Тёзкин", self.make_iin("20080505", "5", "0023")]]})
        rows = {r["full_name"]: r for r in self.upload("/api/students/import/preview", data).get_json()["rows"]}
        self.assertEqual((rows["Без Иина"]["status"], rows["Без Иина"]["match"]), ("update", "name"))
        self.assertEqual(rows["Тёзка Тёзкин"]["status"], "new")
        self.assertTrue(any("другим ИИН" in w for w in rows["Тёзка Тёзкин"]["warnings"]))

    def test_group_spelling_is_unified_with_existing(self):
        from app.models import Student
        from app.extensions import db

        with self.app.app_context():
            db.session.add(Student(full_name="Старый Студент", group_name="ПО 25-1"))
            db.session.commit()
        # Latin "O" in «ПO» and an en dash
        data = self.xlsx({"X": [["ФИО", "ИИН", "Группа"], ["Новый Студент", self.make_iin(), "ПO 25–1"]]})
        row = self.upload("/api/students/import/preview", data).get_json()["rows"][0]
        self.assertEqual(row["group_name"], "ПО 25-1")

    def test_unknown_partner_is_a_warning_not_a_new_partner(self):
        from app.models import Partner

        data = self.xlsx({"ПО 25-1": [["ФИО", "ИИН", "Предприятие"], ["Иванов Иван", self.make_iin(), "Нет Такой ТОО"]]})
        row = self.upload("/api/students/import/preview", data).get_json()["rows"][0]
        self.assertEqual(row["status"], "new")
        self.assertTrue(any("не найдено в реестре" in w for w in row["warnings"]))
        with self.app.app_context():
            self.assertEqual(Partner.query.count(), 0)

    def test_commit_refuses_stale_preview(self):
        data = self.xlsx({"ПО 25-1": [["ФИО", "ИИН"], ["Иванов Иван", self.make_iin()]]})
        r = self.upload("/api/students/import/commit", data, expect_new=5, expect_update=0)
        self.assertEqual(r.status_code, 409)
        self.assertEqual(r.get_json()["code"], "plan_changed")

    # ── RBAC / template / export ───────────────────────────────────────────
    def test_viewer_cannot_import_but_can_download_template(self):
        data = self.xlsx({"A": [["ФИО"], ["Иванов Иван"]]})
        self.assertEqual(self.upload("/api/students/import/preview", data, headers=self.viewer).status_code, 403)
        self.assertEqual(self.upload("/api/students/import/commit", data, headers=self.viewer).status_code, 403)
        r = self.client.get("/api/students/import/template", headers=self.viewer)
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r.mimetype)
        self.assertEqual(self.client.get("/api/students/import/template").status_code, 401)

    def test_template_roundtrips_through_the_importer(self):
        from io import BytesIO
        from openpyxl import load_workbook

        raw = self.client.get("/api/students/import/template", headers=self.admin).data
        wb = load_workbook(BytesIO(raw))
        self.assertEqual(wb.sheetnames[:2], ["Инструкция", "Студенты"])
        ws = wb["Студенты"]
        iin = self.make_iin()
        for col, value in enumerate(["Шаблонов Тест", iin, "ПО 25-1"], start=1):
            ws.cell(row=2, column=col, value=value)
        out = BytesIO()
        wb.save(out)
        body = self.upload("/api/students/import/preview", out.getvalue()).get_json()
        self.assertEqual((body["summary"]["new"], body["summary"]["error"]), (1, 0))

    def test_export_roundtrip(self):
        from app.models import Student
        from app.extensions import db

        iin = self.make_iin()
        with self.app.app_context():
            db.session.add(Student(full_name="Экспорт Тест", iin=iin, group_name="ПО 25-1", course=2,
                                   birth_date=date(2008, 5, 12)))
            db.session.commit()
        raw = self.client.get("/api/students/export?group=ПО 25-1", headers=self.admin).data
        body = self.upload("/api/students/import/preview", raw, {"mode": "fill"}).get_json()
        self.assertEqual(body["rows"][0]["status"], "unchanged")


if __name__ == "__main__":
    import unittest

    unittest.main()
